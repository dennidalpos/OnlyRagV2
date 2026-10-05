"""Prepare and read back retained synthetic translation files without opening app data."""
import hashlib
import json
import os
from pathlib import Path
import sys

import docx
import pymupdf


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")


def prepare(evidence):
    corpus_path = Path(__file__).resolve().parents[2] / "sidecar/tests/fixtures/translation-fidelity-it-en.json"
    corpus = json.loads(corpus_path.read_text(encoding="utf-8"))
    (evidence / "corpus.json").write_bytes(corpus_path.read_bytes())
    inputs = evidence / "inputs"
    inputs.mkdir()
    manifest = {"corpus_sha256": sha256(corpus_path), "cases": []}
    for case in corpus["cases"]:
        source = inputs / (case["id"] + (".docx" if case["format"] == "docx" else ".pdf"))
        if case["format"] == "docx":
            document = docx.Document()
            for paragraph in case["paragraphs"]:
                document.add_paragraph(paragraph)
            document.save(source)
        else:
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                for index, paragraph in enumerate(case["paragraphs"]):
                    bounds = pymupdf.Rect(50, 75 + index * 100, 545, 165 + index * 100)
                    if page.insert_textbox(bounds, paragraph, fontsize=12) < 0:
                        raise ValueError(f"Frozen paragraph does not fit: {case['id']}")
                if case["format"] == "scanned-pdf":
                    image = page.get_pixmap(dpi=200).tobytes("png")
                    with pymupdf.open() as scan:
                        target = scan.new_page()
                        target.insert_image(target.rect, stream=image)
                        scan.save(source)
                else:
                    pdf.save(source)
        manifest["cases"].append({**case, "source": str(source), "source_sha256": sha256(source)})
    write_json(evidence / "manifest.json", manifest)


def readback(evidence):
    report = json.loads((evidence / "report.json").read_text(encoding="utf-8"))
    for case in report["cases"]:
        source = Path(case["source"])
        if sha256(source) != case["source_sha256"]:
            raise ValueError(f"Source changed: {case['id']}")
        output = Path(case["output"])
        case["output_sha256"] = sha256(output)
        if output.suffix == ".docx":
            document = docx.Document(output)
            case["output_text"] = "\n".join(paragraph.text for paragraph in document.paragraphs)
        else:
            with pymupdf.open(output) as pdf:
                case["output_text"] = "\n".join(page.get_text() for page in pdf)
                for index, page in enumerate(pdf):
                    page.get_pixmap(dpi=100).save(evidence / f"{case['id']}-output-{index + 1}.png")
            with pymupdf.open(source) as pdf:
                pdf[0].get_pixmap(dpi=100).save(evidence / f"{case['id']}-source.png")
        if not case["output_text"].strip():
            raise ValueError(f"No output text: {case['id']}")
    write_json(evidence / "report.json", report)


def review_gate(evidence):
    report_path = evidence / "report.json"
    report = json.loads(report_path.read_text(encoding="utf-8"))
    corpus_path = evidence / "corpus.json"
    corpus = json.loads(corpus_path.read_text(encoding="utf-8"))
    review = json.loads((evidence / "review.json").read_text(encoding="utf-8"))
    if not report["complete"] or not review.get("reviewer"):
        raise ValueError("A complete capture and independent reviewer are required.")
    hashes = {"report.json": sha256(report_path), "corpus.json": sha256(corpus_path)}
    if review["artifact_sha256"] != hashes or report["corpus_sha256"] != hashes["corpus.json"]:
        raise ValueError("Review must match the exact captured evidence.")
    expected_ids = {case["id"] for case in corpus["cases"]}
    for collection in (report["cases"], review["cases"]):
        if len(collection) != len(expected_ids) or {case["id"] for case in collection} != expected_ids:
            raise ValueError("Every frozen case requires one capture and one review.")
    for case in report["cases"]:
        if sha256(Path(case["source"])) != case["source_sha256"] or sha256(Path(case["output"])) != case["output_sha256"]:
            raise ValueError("Reviewed source/output bytes changed.")
    for case in review["cases"]:
        if not case.get("notes") or any(type(case.get(key)) is not bool for key in ("content_pass", "language_pass")):
            raise ValueError("Explicit content/language judgments and supporting notes are required.")
    accepted = sum(case["content_pass"] and case["language_pass"] for case in review["cases"])
    result = {"passed": accepted == len(expected_ids), "accepted": accepted, "cases": len(expected_ids),
              "artifact_sha256": hashes, "reviewer": review["reviewer"],
              "limits": "Four synthetic layout-translation cases only; no broad legal, OCR, DOCX visual or model qualification."}
    write_json(evidence / "acceptance.json", result)
    print(json.dumps(result))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    operation, directory = sys.argv[1:]
    target = Path(directory).resolve()
    live_root = Path(os.environ.get("ONLYRAG_LIVE_ROOT", Path.home() / "OnlyRag-Live")).resolve()
    if target.parent != live_root or not target.name.startswith("translation-fidelity-"):
        raise ValueError("Use a retained translation campaign directly inside OnlyRag-Live.")
    if operation == "prepare":
        prepare(target)
    elif operation == "readback":
        readback(target)
    elif operation == "review":
        raise SystemExit(review_gate(target))
    else:
        raise ValueError("Expected prepare, readback or review.")
