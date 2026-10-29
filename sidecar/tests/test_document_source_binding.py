"""Retained source consumers exercised with native containers and isolated LanceDB."""
import base64
import io
import json
from pathlib import Path

import docx
import pymupdf
import pytest
from PIL import Image

from sidecar.domain import translator
from sidecar.infrastructure import db, source_provenance as source
from sidecar.services import ingest_service, search_service


def files(directory):
    return {str(p.relative_to(directory)): p.read_bytes() for p in directory.rglob("*") if p.is_file()}


def write_source(path, text="ORIGINAL_CANONICAL_TEXT"):
    if path.suffix == ".docx":
        document = docx.Document()
        document.add_paragraph(text)
        document.save(path)
    elif path.suffix == ".pdf":
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), text)
            document.save(path)
    else:
        Image.new("RGB", (12, 9), "red" if text.startswith("ORIGINAL") else "blue").save(path)


@pytest.fixture(params=["docx", "pdf", "png"])
def selected(request, tmp_path, monkeypatch):
    monkeypatch.setattr(source, "SOURCE_ROOT", tmp_path / "archives")
    monkeypatch.setattr("sidecar.domain.ingestion.run_page_ocr", lambda *args, **kwargs: "ORIGINAL_CANONICAL_TEXT")
    path = tmp_path / f"source.{request.param}"
    write_source(path)
    content = path.read_bytes()
    events = [json.loads(line) for line in ingest_service.process_and_index_document_generator(str(path))]
    assert events[-1]["type"] == "done", events
    record = search_service.get_stored_document(events[-1]["data"]["id"])
    yield path, record, content


def translate(record, tmp_path, monkeypatch):
    seen = []
    def runs(items, *args):
        seen.extend(item.text for item in items)
        for item in items:
            item.text = "Translated text"
    def blocks(items, *args):
        seen.extend(item["text"] for item in items)
        for item in items:
            item["text"] = "Translated text"
    monkeypatch.setattr(translator, "_translate_batch", runs)
    monkeypatch.setattr(translator, "_translate_pdf_blocks", blocks)
    events = [json.loads(line) for line in translator.translate_document_stream(
        record, "English", "Italian", "deterministic", target_dir=str(tmp_path / "outputs"))]
    return events, seen


@pytest.mark.parametrize("change", ["modify", "move", "remove"])
def test_callers_cannot_replace_translation_or_preview_bytes(selected, tmp_path, monkeypatch, sidecar_http_client, change):
    path, record, original = selected
    before = ingest_service.render_document_page_preview(record["id"], 1)
    archives = files(source.SOURCE_ROOT)
    if change == "modify":
        write_source(path, "CHANGED_EXTERNAL_TEXT")
    elif change == "move":
        path.rename(path.with_name("moved" + path.suffix))
    else:
        path.unlink()
    after = ingest_service.render_document_page_preview(record["id"], 1)
    assert after == before
    response = sidecar_http_client.get(f"/documents/{record['id']}/page-preview/1")
    assert response.status_code == 200 and response.json() == after.model_dump()
    image = Image.open(io.BytesIO(base64.b64decode(after.image_base64)))
    assert image.format == "PNG"
    if path.suffix != ".png":
        prepared = translator.prepare_translation(record["id"])
        events, seen = translate(prepared, tmp_path, monkeypatch)
        assert events[-1]["type"] == "done", events
        assert seen and all("CHANGED_EXTERNAL_TEXT" not in text for text in seen)
        assert any("ORIGINAL_CANONICAL_TEXT" in text for text in seen)
        output = Path(events[-1]["data"]["filePath"])
        if path.suffix == ".docx":
            assert docx.Document(output).paragraphs[0].text == "Translated text"
        else:
            with pymupdf.open(output) as document:
                assert "Translated text" in document[0].get_text()
    assert files(source.SOURCE_ROOT) == archives
    assert search_service.get_stored_document(record["id"]) == record


@pytest.mark.parametrize("fault", ["missing_original", "changed_original", "missing_archive", "changed_archive", "read_failure"])
def test_missing_corrupt_or_unreadable_evidence_is_refused(selected, tmp_path, monkeypatch, sidecar_http_client, fault):
    path, record, original = selected
    directory = source.SOURCE_ROOT / record["id"]
    retained = directory / f"original{path.suffix}"
    metadata = directory / "extraction.json"
    if fault == "missing_original":
        retained.unlink()
    elif fault == "changed_original":
        retained.write_bytes(b"Different retained bytes")
    elif fault == "missing_archive":
        metadata.unlink()
    elif fault == "changed_archive":
        data = json.loads(metadata.read_text(encoding="utf-8"))
        data["raw_markdown"] += "Corrupt extraction"
        metadata.write_text(json.dumps(data), encoding="utf-8")
    else:
        native_open = Path.open
        def refuse(target, *args, **kwargs):
            if target == retained:
                raise PermissionError("Synthetic source read denial")
            return native_open(target, *args, **kwargs)
        monkeypatch.setattr(Path, "open", refuse)
    before = files(source.SOURCE_ROOT) if fault != "read_failure" else None
    with pytest.raises(ValueError):
        ingest_service.render_document_page_preview(record["id"], 1)
    assert sidecar_http_client.get(f"/documents/{record['id']}/page-preview/1").status_code == 404
    if path.suffix != ".png":
        with pytest.raises(ValueError):
            translator.prepare_translation(record["id"])
        assert sidecar_http_client.post(f"/documents/{record['id']}/translate-inplace-stream", json={
            "source_lang": "English", "target_lang": "Italian", "model": "deterministic",
        }).status_code == 404
        events, seen = translate(record, tmp_path, monkeypatch)
        assert events[-1]["type"] == "error" and seen == []
        assert not (tmp_path / "outputs").exists()
    if before is not None:
        assert files(source.SOURCE_ROOT) == before
    assert path.read_bytes() == original
    assert search_service.get_stored_document(record["id"]) == record


@pytest.mark.parametrize("phase", ["before_stream", "during_translation", "delete_during_translation"])
@pytest.mark.parametrize("selected", ["docx", "pdf"], indirect=True)
def test_changed_canonical_revision_never_publishes_stale_translation(selected, tmp_path, monkeypatch, phase):
    path, record, _ = selected
    prepared = translator.prepare_translation(record["id"])
    archives = files(source.SOURCE_ROOT)
    def change(*args):
        if phase == "delete_during_translation":
            db.delete_document_records(record["id"])
        else:
            ingest_service.update_and_reindex_document(record["id"], "Revised canonical text.", record["ingested_at"])
    if phase == "before_stream":
        change()
    else:
        monkeypatch.setattr(translator, "_translate_batch", change)
        monkeypatch.setattr(translator, "_translate_pdf_blocks", change)
    events = [json.loads(line) for line in translator.translate_document_stream(
        prepared, "English", "Italian", "deterministic", target_dir=str(tmp_path / "outputs"))]
    assert events[-1]["type"] == "error", events
    assert "revision" in events[-1]["error"]
    assert not list((tmp_path / "outputs").glob("*"))
    assert files(source.SOURCE_ROOT) == archives


def test_preview_of_original_stays_bound_after_canonical_edit(selected):
    _, record, _ = selected
    before = ingest_service.render_document_page_preview(record["id"], 1)
    ingest_service.update_and_reindex_document(record["id"], "Revised canonical text.", record["ingested_at"])
    assert ingest_service.render_document_page_preview(record["id"], 1) == before


def test_preexisting_record_without_archive_refuses_without_migrating(tmp_path, monkeypatch):
    monkeypatch.setattr(source, "SOURCE_ROOT", tmp_path / "archives")
    path = tmp_path / "legacy.docx"
    write_source(path)
    record = dict(id="legacy", file_path=str(path), filename=path.name, file_type="docx", ingested_at="old",
                  extracted_markdown="Legacy canonical text", status="indexed", num_chunks=0, num_pages=1,
                  file_size=path.stat().st_size, used_fallback_embeddings=False)
    db.write_document_records(record, [])
    with pytest.raises(ValueError, match="archive"):
        translator.prepare_translation("legacy")
    with pytest.raises(ValueError, match="archive"):
        ingest_service.render_document_page_preview("legacy", 1)
    assert not source.SOURCE_ROOT.exists()
    assert search_service.get_stored_document("legacy") == record


def test_archive_revision_is_rechecked_between_preflight_and_dispatch(tmp_path, monkeypatch):
    from _stream import retained_translation_record
    path = tmp_path / "source.docx"
    write_source(path)
    prepared = retained_translation_record(path)
    directory = source.SOURCE_ROOT / prepared["id"]
    write_source(directory / "original.docx", "CHANGED_EXTERNAL_TEXT")
    import hashlib
    metadata = directory / "extraction.json"
    data = json.loads(metadata.read_text(encoding="utf-8"))
    data["source_revision"] = hashlib.sha256((directory / "original.docx").read_bytes()).hexdigest()
    metadata.write_text(json.dumps(data), encoding="utf-8")
    events, seen = translate(prepared, tmp_path, monkeypatch)
    assert events[-1]["type"] == "error" and "revision" in events[-1]["error"]
    assert seen == [] and not (tmp_path / "outputs").exists()


def test_verified_bytes_are_consumed_even_if_archive_path_changes_after_read(tmp_path, monkeypatch):
    from _stream import retained_translation_record
    path = tmp_path / "source.docx"
    write_source(path)
    prepared = retained_translation_record(path)
    native_read = source.read_retained_source
    def replace_after_read(record):
        archive, content = native_read(record)
        write_source(source.SOURCE_ROOT / record["id"] / archive.original_file, "CHANGED_EXTERNAL_TEXT")
        return archive, content
    monkeypatch.setattr(source, "read_retained_source", replace_after_read)
    events, seen = translate(prepared, tmp_path, monkeypatch)
    assert events[-1]["type"] == "done", events
    assert seen == ["ORIGINAL_CANONICAL_TEXT"]
    assert docx.Document(events[-1]["data"]["filePath"]).paragraphs[0].text == "Translated text"


def test_retained_read_budget_refuses_before_consumers(selected, monkeypatch):
    _, record, _ = selected
    monkeypatch.setattr(source, "MAX_SOURCE_BYTES", 8)
    with pytest.raises(ValueError, match="budget"):
        ingest_service.render_document_page_preview(record["id"], 1)
