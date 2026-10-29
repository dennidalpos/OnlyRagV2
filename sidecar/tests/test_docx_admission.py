"""Native DOCX admission and NDJSON/LanceDB boundaries, with no model calls."""
import io
import json
from pathlib import Path
from unittest.mock import Mock

import docx
import pytest

from sidecar.domain.ingestion import extract_document_markdown
from sidecar.infrastructure import db, source_provenance as source
from sidecar.services import ingest_service


def rows():
    return {name: db.lance_db.open_table(name).to_arrow().to_pylist() for name in db.get_existing_tables()}


def files(directory):
    return {str(p.relative_to(directory)): p.read_bytes() for p in directory.rglob("*") if p.is_file()}


@pytest.fixture(params=["malformed", "truncated", "encrypted", "unreadable", "empty"])
def refused_docx(request, tmp_path, monkeypatch):
    document = docx.Document()
    if request.param != "empty":
        document.add_paragraph("Retain the original Word text.")
    stream = io.BytesIO()
    document.save(stream)
    content = stream.getvalue()
    if request.param == "malformed":
        content = b"Not a ZIP package. Never index this fallback text."
    elif request.param == "truncated":
        content = content[:len(content) // 2]
    elif request.param == "encrypted":
        content = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 32 + b"EncryptedPackage EncryptionInfo"
    file = tmp_path / f"{request.param}.docx"
    file.write_bytes(content)
    if request.param == "unreadable":
        original = docx.Document
        def refuse_read(value=None):
            if value is not None:
                raise PermissionError("Synthetic DOCX read failure")
            return original()
        monkeypatch.setattr(docx, "Document", refuse_read)
    return file, content


def test_extraction_refuses_docx_without_binary_fallback(refused_docx):
    file, content = refused_docx
    for path in (None, str(file)):
        with pytest.raises(ValueError, match="DOCX"):
            extract_document_markdown(file.name, content, path)
    assert file.read_bytes() == content


def test_stream_refuses_docx_before_chunk_embedding_or_commit(refused_docx, tmp_path, monkeypatch):
    monkeypatch.setattr(source, "SOURCE_ROOT", tmp_path / "archives")
    file, content = refused_docx
    old = tmp_path / "existing.md"
    old.write_text("Preserve the existing document and its chunks.", encoding="utf-8")
    assert json.loads(list(ingest_service.process_and_index_document_generator(str(old)))[-1])["type"] == "done"
    before, archives = rows(), files(source.SOURCE_ROOT)
    for name in ("create_semantic_chunks", "generate_embedding_batch", "write_document_records"):
        monkeypatch.setattr(ingest_service, name, Mock(side_effect=AssertionError(f"Refused DOCX reached {name}")))
    events = [json.loads(line) for line in ingest_service.process_and_index_document_generator(str(file))]
    assert events[-1]["type"] == "error"
    assert "DOCX" in events[-1]["error"]
    assert not any(event["type"] == "done" or event.get("step_code") in ("chunking", "embedding") for event in events)
    for name in ("create_semantic_chunks", "generate_embedding_batch", "write_document_records"):
        getattr(ingest_service, name).assert_not_called()
    assert rows() == before
    after = files(source.SOURCE_ROOT)
    assert all(after[path] == value for path, value in archives.items())
    added = set(after) - set(archives)
    assert len(added) == 1 and Path(next(iter(added))).name == "original.docx"
    assert after[next(iter(added))] == content
    assert file.read_bytes() == content


def test_valid_docx_still_extracts_indexes_and_retains_exact_bytes(tmp_path, monkeypatch):
    monkeypatch.setattr(source, "SOURCE_ROOT", tmp_path / "archives")
    file = tmp_path / "valid.docx"
    document = docx.Document()
    document.add_paragraph("Valid Word content must remain searchable.")
    document.save(file)
    before = file.read_bytes()
    events = [json.loads(line) for line in ingest_service.process_and_index_document_generator(str(file))]
    assert events[-1]["type"] == "done", events
    data = events[-1]["data"]
    assert "Valid Word content" in data["extracted_markdown"]
    assert data["num_chunks"] > 0
    assert (source.SOURCE_ROOT / data["id"] / "original.docx").read_bytes() == before == file.read_bytes()
    assert len(rows()["documents"]) == 1 and len(rows()["chunks"]) == data["num_chunks"]
