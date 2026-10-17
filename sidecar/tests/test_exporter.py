import base64
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

import docx
import pytest

from sidecar.domain import exporter


@pytest.fixture
def export_root(tmp_path, monkeypatch):
    monkeypatch.setattr(exporter, "EXPORT_DIR", str(tmp_path))
    return tmp_path


class ParsedHtml(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags = []
        self.in_pre = False
        self.text = []

    def handle_starttag(self, tag, attrs):
        self.tags.append(tag)
        if tag == "pre":
            self.in_pre = True

    def handle_endtag(self, tag):
        if tag == "pre":
            self.in_pre = False

    def handle_data(self, data):
        if self.in_pre:
            self.text.append(data)


def test_html_export_keeps_untrusted_markup_as_literal_unicode_text(export_root):
    content = '</pre><script>alert("audit")</script><img src="https://example.invalid/pixel" onerror="alert(1)"> &amp; città 漢 😀'
    result = exporter.export_markdown_to_file(content, "html")
    parsed = ParsedHtml()
    parsed.feed(base64.b64decode(result["base64_content"]).decode("utf-8"))
    assert not {"script", "img", "iframe"}.intersection(parsed.tags)
    assert "".join(parsed.text) == content
    assert Path(result["file_path"]).read_bytes() == base64.b64decode(result["base64_content"])


@pytest.mark.parametrize("fmt", ["pdf", "docx", "html", "htm"])
def test_same_instant_exports_preserve_both_files(export_root, monkeypatch, fmt):
    class FrozenDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 10, 17, 12, 0, 0)

    monkeypatch.setattr(exporter.datetime, "datetime", FrozenDateTime)
    first = exporter.export_markdown_to_file("# First document", fmt)
    original = Path(first["file_path"]).read_bytes()
    second = exporter.export_markdown_to_file("# Second document", fmt)
    assert first["file_path"] != second["file_path"]
    assert Path(first["file_path"]).read_bytes() == original
    for result in (first, second):
        assert Path(result["file_path"]).read_bytes() == base64.b64decode(result["base64_content"])


def test_concurrent_html_exports_keep_each_own_content(export_root):
    contents = [f"Document {i}: città 漢 😀" for i in range(12)]
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda content: exporter.export_markdown_to_file(content, "html"), contents))
    assert len({result["file_path"] for result in results}) == len(contents)
    for content, result in zip(contents, results):
        parsed = ParsedHtml()
        parsed.feed(Path(result["file_path"]).read_text(encoding="utf-8"))
        assert "".join(parsed.text) == content


def test_failed_render_preserves_previous_export_and_reports_failure(export_root, monkeypatch):
    previous = exporter.export_markdown_to_file("Preserve this document", "docx")
    original = Path(previous["file_path"]).read_bytes()

    def fail_after_partial_write(content, output_path):
        Path(output_path).write_bytes(b"incomplete")
        raise OSError("Disk full")

    monkeypatch.setattr(exporter, "_render_docx_from_markdown", fail_after_partial_write)
    with pytest.raises(RuntimeError, match="DOCX export failed"):
        exporter.export_markdown_to_file("New document", "docx")
    assert Path(previous["file_path"]).read_bytes() == original


@pytest.mark.parametrize("closing", ["", "\n```"])
@pytest.mark.parametrize("prefix", ["before\n", "| Column |\n| --- |\n| retained cell |\n"])
def test_docx_keeps_code_at_end_and_after_tables(export_root, closing, prefix):
    result = exporter.export_markdown_to_file(prefix + "```python\nvaluable_code = 'città 漢 😀'" + closing, "docx")
    document = docx.Document(result["file_path"])
    assert [p.text for p in document.paragraphs][-1] == "valuable_code = 'città 漢 😀'"
    if prefix.startswith("|"):
        assert document.tables[0].cell(1, 0).text == "retained cell"
        body = list(document.element.body)
        assert body.index(document.tables[0]._element) < body.index(document.paragraphs[-1]._element)
    else:
        assert document.paragraphs[0].text == "before"


@pytest.mark.parametrize("content", ["before\n```", "before\n```\n```", "before\n```\n```\nafter"])
def test_docx_empty_code_blocks_preserve_surrounding_text(export_root, content):
    result = exporter.export_markdown_to_file(content, "docx")
    texts = [p.text for p in docx.Document(result["file_path"]).paragraphs]
    assert texts[0] == "before"
    if content.endswith("after"):
        assert texts[-1] == "after"
