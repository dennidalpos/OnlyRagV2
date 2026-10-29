"""Native DOCX/TIFF fixtures, complete extraction and isolated NDJSON/LanceDB readback."""
from collections import Counter
import io
import json
import random
import xml.etree.ElementTree as ET
from zipfile import ZipFile
from unittest.mock import Mock

import docx
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.shared import Inches
from PIL import Image
import pytest

from sidecar.domain import ingestion
from sidecar.domain.router import classify_file_type, DocumentCategory
from sidecar.infrastructure import db, source_provenance as source
from sidecar.services import ingest_service, search_service


@pytest.fixture(autouse=True)
def isolated_archives(tmp_path, monkeypatch):
    monkeypatch.setattr(source, "SOURCE_ROOT", tmp_path / "archives")


def files(root):
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in root.rglob("*") if p.is_file()}


def rows():
    return {name: sorted(db.lance_db.open_table(name).to_arrow().to_pylist(),
                         key=lambda row: row.get("id", row.get("chunk_id", ""))) for name in db.get_existing_tables()}


def ingest(path):
    return [json.loads(line) for line in ingest_service.process_and_index_document_generator(str(path))]


def survivor(tmp_path):
    path = tmp_path / "existing.md"
    path.write_text("Preserve existing canonical text and original bytes.", encoding="utf-8")
    assert ingest(path)[-1]["type"] == "done"
    return rows(), files(source.SOURCE_ROOT)


def assert_refused(path, monkeypatch, before, archives):
    for name in ("create_semantic_chunks", "generate_embedding_batch", "write_document_records"):
        monkeypatch.setattr(ingest_service, name, Mock(side_effect=AssertionError(f"Refused input reached {name}")))
    original = path.read_bytes()
    events = ingest(path)
    assert events[-1]["type"] == "error", events
    assert not any(event["type"] == "done" for event in events)
    for name in ("create_semantic_chunks", "generate_embedding_batch", "write_document_records"):
        getattr(ingest_service, name).assert_not_called()
    assert rows() == before
    after = files(source.SOURCE_ROOT)
    assert all(after[name] == data for name, data in archives.items())
    added = set(after) - set(archives)
    assert len(added) == 1 and "/original." in next(iter(added))
    assert after[next(iter(added))] == original == path.read_bytes()
    return events[-1]["error"]


@pytest.mark.parametrize("variant", ["linked", "distinct", "first_even", "no_headers"])
def test_docx_complete_markdown_order_and_shared_story_identity(tmp_path, variant):
    path = tmp_path / "stories.docx"
    doc = docx.Document()
    doc.add_paragraph("BEFORE_TABLE")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "TABLE_MIDDLE"
    doc.tables[0].cell(0, 1).text = "TABLE_RIGHT"
    doc.add_paragraph("AFTER_TABLE")
    headers, footers = [], []
    if variant != "no_headers":
        doc.sections[0].header.paragraphs[0].text = "HEADER_REQUIRED"
        doc.sections[0].footer.paragraphs[0].text = "FOOTER_REQUIRED"
        doc.sections[0].header.add_table(rows=1, cols=1, width=Inches(3)).cell(0, 0).text = "HEADER_TABLE"
        doc.sections[0].footer.add_table(rows=1, cols=1, width=Inches(3)).cell(0, 0).text = "FOOTER_TABLE"
        headers.extend(["HEADER_REQUIRED", "| HEADER_TABLE |\n| --- |"])
        footers.extend(["FOOTER_REQUIRED", "| FOOTER_TABLE |\n| --- |"])
    if variant == "first_even":
        doc.settings.odd_and_even_pages_header_footer = True
        doc.sections[0].different_first_page_header_footer = True
        for name in ("first_page_header", "even_page_header", "first_page_footer", "even_page_footer"):
            getattr(doc.sections[0], name).paragraphs[0].text = name.upper()
            (headers if name.endswith("header") else footers).append(name.upper())
    second = doc.add_section(WD_SECTION_START.NEW_PAGE)
    if variant == "distinct":
        second.header.is_linked_to_previous = False
        second.header.paragraphs[0].text = "SECOND_HEADER"
        second.footer.is_linked_to_previous = False
        second.footer.paragraphs[0].text = "SECOND_FOOTER"
        headers.append("SECOND_HEADER")
        footers.append("SECOND_FOOTER")
    else:
        assert second.header.is_linked_to_previous and second.footer.is_linked_to_previous
    doc.save(path)
    original = path.read_bytes()
    body = ["BEFORE_TABLE", "| TABLE_MIDDLE | TABLE_RIGHT |\n| --- | --- |", "AFTER_TABLE"]
    expected = "# stories.docx\n\n## Page 1\n\n" + "\n\n".join(headers + body + footers)
    for source_path in (None, str(path)):
        markdown, pages = ingestion.extract_document_markdown(path.name, original, source_path)
        assert markdown == expected and pages == 1
    events = ingest(path)
    assert events[-1]["type"] == "done", events
    record = search_service.get_stored_document(events[-1]["data"]["id"])
    assert record["extracted_markdown"] == expected
    with ZipFile(path) as package:
        visible = [node.text for name in package.namelist()
                   if name == "word/document.xml" or name.startswith(("word/header", "word/footer")) and name.endswith(".xml")
                   for node in ET.fromstring(package.read(name)).iter(qn("w:t")) if node.text]
    assert all(expected.count(text) == count for text, count in Counter(visible).items())
    assert source._archive(record["id"]).raw_markdown == expected
    assert (source.SOURCE_ROOT / record["id"] / "original.docx").read_bytes() == original == path.read_bytes()


def test_docx_nested_merged_cells_hyperlinks_headings_and_unicode(tmp_path):
    path = tmp_path / "nested.docx"
    doc = docx.Document()
    doc.add_heading("Titolo Ω 😀", level=2)
    paragraph = doc.add_paragraph("Link: ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "bookmark")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "HYPERLINK_TEXT"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    table = doc.add_table(rows=2, cols=2)
    merged = table.cell(0, 0).merge(table.cell(1, 0))
    merged.text = "MERGED_ONCE"
    nested = merged.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "NESTED_TEXT"
    table.cell(0, 1).text = "A|B\nSECOND_LINE"
    table.cell(1, 1).text = "LAST_CELL"
    doc.add_paragraph("BODY_END")
    doc.save(path)
    expected = ("# nested.docx\n\n## Page 1\n\n## Titolo Ω 😀\n\nLink: HYPERLINK_TEXT\n\n"
                "| MERGED_ONCE<br><br>\\| NESTED_TEXT \\|<br>\\| --- \\| | A\\|B\nSECOND_LINE |\n"
                "| --- | --- |\n|  | LAST_CELL |\n\nBODY_END")
    # Cell line breaks are explicit HTML separators inside the Markdown table.
    expected = expected.replace("A\\|B\nSECOND_LINE", "A\\|B<br>SECOND_LINE")
    events = ingest(path)
    assert events[-1]["type"] == "done", events
    assert events[-1]["data"]["extracted_markdown"] == expected


def test_reading_docx_does_not_materialize_header_footer_parts():
    from sidecar.domain.docx_content import extract_docx_blocks
    doc = docx.Document()
    doc.add_paragraph("Body only")
    before = dict(doc.part.rels)
    assert extract_docx_blocks(doc) == ["Body only"]
    assert dict(doc.part.rels) == before


@pytest.mark.parametrize("tag", ["txbxContent", "altChunk", "fldSimple", "fldChar", "instrText",
                                  "footnoteReference", "endnoteReference", "commentReference", "sdt",
                                  "sym", "ins", "del", "moveFrom", "moveTo"])
def test_unsupported_visible_docx_content_refuses_before_indexing(tmp_path, monkeypatch, tag):
    before, archives = survivor(tmp_path)
    path = tmp_path / "unsupported.docx"
    doc = docx.Document()
    doc.add_paragraph("Do not index partial body text")
    doc.sections[0].header.paragraphs[0]._p.append(OxmlElement(f"w:{tag}"))
    doc.save(path)
    assert "Unsupported DOCX" in assert_refused(path, monkeypatch, before, archives)


def test_additional_visible_word_part_is_explicitly_refused(tmp_path, monkeypatch):
    before, archives = survivor(tmp_path)
    path = tmp_path / "footnotes.docx"
    doc = docx.Document()
    doc.add_paragraph("Body text")
    part = XmlPart(PackURI("/word/footnotes.xml"), "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                   parse_xml('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>VISIBLE_NOTE</w:t></w:r></w:p></w:footnote></w:footnotes>'), doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)
    doc.save(path)
    assert "Unsupported DOCX text part" in assert_refused(path, monkeypatch, before, archives)


@pytest.mark.parametrize("placement", ["inline", "block", "table"])
def test_unknown_word_text_wrappers_are_refused_instead_of_omitted(tmp_path, monkeypatch, placement):
    before, archives = survivor(tmp_path)
    path = tmp_path / "wrapped.docx"
    doc = docx.Document()
    paragraph = doc.add_paragraph("VISIBLE_WRAPPED_TEXT")
    wrapper = OxmlElement("w:customXml")
    if placement == "inline":
        wrapper.append(paragraph.runs[0]._r)
        paragraph._p.append(wrapper)
    elif placement == "block":
        wrapper.append(paragraph._p)
        doc.element.body.insert(0, wrapper)
    else:
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "VISIBLE_TABLE_TEXT"
        wrapper.append(table._tbl.tr_lst[0])
        table._tbl.append(wrapper)
    doc.save(path)
    assert "Unsupported DOCX" in assert_refused(path, monkeypatch, before, archives)


def test_docx_shared_image_ocr_is_preserved_once_and_failure_is_explicit(tmp_path, monkeypatch):
    before, archives = survivor(tmp_path)
    path = tmp_path / "images.docx"
    image = Image.frombytes("RGB", (100, 100), random.Random(8).randbytes(30000))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    doc = docx.Document()
    doc.add_paragraph("BODY_IMAGE")
    doc.add_picture(io.BytesIO(buffer.getvalue()))
    doc.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(buffer.getvalue()))
    doc.save(path)
    calls = []
    monkeypatch.setattr(ingestion, "run_page_ocr", lambda data, **kwargs: calls.append(data) or "OCR_VISIBLE_TEXT")
    events = ingest(path)
    assert events[-1]["type"] == "done", events
    assert calls == [buffer.getvalue()]
    assert events[-1]["data"]["extracted_markdown"] == ("# images.docx\n\n## Page 1\n\nBODY_IMAGE\n\n"
            "### Immagine DOCX 1 (Testo rilevato da OCR)\n\nOCR_VISIBLE_TEXT")
    before, archives = rows(), files(source.SOURCE_ROOT)
    def fail(*args, **kwargs):
        raise ValueError("Controlled OCR failure")
    monkeypatch.setattr(ingestion, "run_page_ocr", fail)
    assert "image extraction refused" in assert_refused(path, monkeypatch, before, archives)


@pytest.mark.parametrize("suffix", ["tif", "tiff", "TIF", "TIFF", "TiF", "TiFf", "png", "jpg", "jpeg", "bmp", "gif", "webp"])
def test_image_aliases_use_native_container_and_controlled_ocr_through_ingestion(tmp_path, monkeypatch, suffix):
    before, archives = survivor(tmp_path)
    path = tmp_path / f"scan.{suffix}"
    format = "TIFF" if suffix.lower() in ("tif", "tiff") else ("JPEG" if suffix in ("jpg", "jpeg") else suffix.upper())
    image = Image.new("RGB", (30, 20), "white")
    buffer = io.BytesIO()
    image.save(buffer, format=format)
    payload = buffer.getvalue()
    path.write_bytes(payload)
    assert classify_file_type(path.name) == DocumentCategory.IMAGE
    calls = []
    monkeypatch.setattr(ingestion, "run_page_ocr", lambda data, **kwargs: calls.append(data) or "NATIVE_IMAGE_TEXT")
    markdown, pages = ingestion.extract_document_markdown(path.name, payload)
    assert markdown == f"# {path.name}\n\n## Page 1\n\nNATIVE_IMAGE_TEXT" and pages == 1
    events = ingest(path)
    assert events[-1]["type"] == "done", events
    data = events[-1]["data"]
    assert data["extracted_markdown"] == markdown
    assert calls == [payload, payload]
    assert search_service.get_stored_document(data["id"])["extracted_markdown"] == markdown
    assert source._archive(data["id"]).raw_markdown == markdown
    assert (source.SOURCE_ROOT / data["id"] / f"original.{suffix.lower()}").read_bytes() == payload == path.read_bytes()
    after = rows()
    for table, key in (("documents", "id"), ("chunks", "doc_id")):
        assert [row for row in after[table] if row[key] != data["id"]] == before[table]
    assert all(files(source.SOURCE_ROOT)[name] == value for name, value in archives.items())


@pytest.mark.parametrize("suffix", ["tif", "tiff", "TIF", "TIFF"])
@pytest.mark.parametrize("damage", ["header", "truncated_native"])
def test_malformed_tiff_never_dispatches_ocr_or_indexes_container_text(tmp_path, monkeypatch, suffix, damage):
    before, archives = survivor(tmp_path)
    path = tmp_path / f"broken.{suffix}"
    if damage == "header":
        path.write_bytes(b"II*\x00" + b"Unusable truncated TIFF directory")
    else:
        image = io.BytesIO()
        Image.new("RGB", (30, 20), "white").save(image, format="TIFF")
        path.write_bytes(image.getvalue()[:len(image.getvalue()) // 2])
    dispatch = Mock(side_effect=AssertionError("Malformed TIFF reached OCR"))
    monkeypatch.setattr(ingestion, "run_page_ocr", dispatch)
    assert "Image container" in assert_refused(path, monkeypatch, before, archives)
    dispatch.assert_not_called()


def test_identical_tiff_bytes_have_identical_content_and_dispatch_across_aliases(tmp_path, monkeypatch):
    payload = io.BytesIO()
    Image.new("RGB", (30, 20), "white").save(payload, format="TIFF")
    original = payload.getvalue()
    calls = []
    monkeypatch.setattr(ingestion, "run_page_ocr", lambda data, **kwargs: calls.append(data) or "SAME_TIFF_TEXT")
    for suffix in ("tif", "tiff", "TIF", "TIFF"):
        path = tmp_path / f"identical.{suffix}"
        path.write_bytes(original)
        events = ingest(path)
        assert events[-1]["type"] == "done", events
        data = events[-1]["data"]
        assert data["extracted_markdown"] == f"# {path.name}\n\n## Page 1\n\nSAME_TIFF_TEXT"
        archive = source._archive(data["id"])
        assert (source.SOURCE_ROOT / data["id"] / archive.original_file).read_bytes() == original
        assert path.read_bytes() == original
    assert calls == [original] * 4
    assert len(rows()["documents"]) == 4
