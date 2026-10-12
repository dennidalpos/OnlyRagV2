import io
import sys
import zipfile
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from sidecar.domain.ingestion import extract_document_markdown
from _stream import done_payload, read_events


FIXTURES = Path(__file__).parent / "fixtures"


@pytest.mark.parametrize("extension,separator", [("csv", ","), ("tsv", "\t")])
@pytest.mark.parametrize("from_disk", [False, True])
def test_delimited_cells_preserve_literal_values(tmp_path, extension, separator, from_disk):
    content = separator.join(["Code", "Status", "Amount", "Note"]) + "\n"
    content += separator.join(["0007", "NA", "12.50", '"Ada, Neri"']) + "\n"
    content += separator.join(["0008", "NULL", "0", '"first line\nsecond line"']) + "\n"
    original = content.encode("utf-8")
    path = tmp_path / f"literal.{extension}"
    path.write_bytes(original)
    markdown, pages = extract_document_markdown(path.name, b"" if from_disk else original, str(path) if from_disk else None)
    assert pages == 1
    for value in ("0007", "NA", "12.50", "Ada, Neri", "0008", "NULL", "first line", "second line"):
        assert value in markdown
    assert "nan" not in markdown
    assert path.read_bytes() == original


@pytest.mark.parametrize("content", [
    b'Code,Status\n0007,"unterminated\n',
    b"Code,Status\n0007,Ready\n0008,Ready,unexpected\n",
    b"Code,Status\n0007,Ready,unexpected\n0008,Ready,unexpected\n",
])
def test_malformed_csv_never_indexes_partial_content(sidecar_http_client, tmp_path, content):
    path = tmp_path / "malformed.csv"
    path.write_bytes(content)
    before = sidecar_http_client.get("/health").json()
    events = read_events(sidecar_http_client.post("/ingest-path-stream", json={"file_path": str(path)}))
    assert events[-1]["type"] == "error"
    assert not any(event["type"] == "done" or event.get("step_code") == "embedding" for event in events)
    after = sidecar_http_client.get("/health").json()
    assert (after["documents_count"], after["chunks_count"]) == (before["documents_count"], before["chunks_count"])
    assert path.read_bytes() == content


def test_csv_stream_preserves_the_original_literals(sidecar_http_client, tmp_path):
    content = b"Order,Route,Status\n0007,North,NA\n0008,South,Ready\n"
    path = tmp_path / "literal.csv"
    path.write_bytes(content)
    payload = done_payload(sidecar_http_client.post("/ingest-path-stream", json={"file_path": str(path)}))
    for value in ("0007", "NA", "0008", "Ready"):
        assert value in payload["extracted_markdown"]
    assert path.read_bytes() == content


def test_csv_late_malformed_row_after_display_limit_fails():
    content = "Code,Status\n" + "0007,Ready\n" * 1100 + "0008,Ready,unexpected\n"
    with pytest.raises((pd.errors.ParserError, pd.errors.ParserWarning)):
        extract_document_markdown("late.csv", content.encode(), max_tabular_rows=2)


def test_csv_counts_omitted_rows_across_chunks():
    content = "Code,Status\n" + "".join(f"{index},Ready\n" for index in range(1100))
    markdown, _ = extract_document_markdown("many.csv", content.encode(), max_tabular_rows=2)
    assert "su 1100 totali" in markdown
    assert "1099" not in markdown


def test_csv_blank_data_row_is_counted_instead_of_skipped():
    markdown, _ = extract_document_markdown("blank.csv", b"Code,Status\nfirst,Ready\n\nlast,Ready\n", max_tabular_rows=1)
    assert "su 3 totali" in markdown
    assert "last" not in markdown


def test_csv_refuses_oversize_field_and_wide_record_without_changing_source(tmp_path):
    cases = [
        (b"Key,Value\n1," + b"x" * (64 * 1024 + 1) + b"\n", "cell exceeds"),
        (b",".join(f"C{i}".encode() for i in range(257)) + b"\n", "column parser limit"),
    ]
    for index, (content, message) in enumerate(cases):
        source = tmp_path / f"adversarial-{index}.csv"
        source.write_bytes(content)
        with pytest.raises(ValueError, match=message):
            extract_document_markdown(source.name, b"", str(source))
        assert source.read_bytes() == content


def test_csv_selected_output_budget_refuses_without_truncation():
    content = b"Value\n" + (b"x" * (64 * 1024) + b"\n") * 65
    with pytest.raises(ValueError, match="render limit"):
        extract_document_markdown("large.csv", content)


@pytest.mark.parametrize("extension", ["csv", "xlsx", "parquet"])
def test_adversarial_table_stream_never_indexes_partial_content(sidecar_http_client, tmp_path, extension):
    if extension == "csv":
        content = b"Code,Value\n0007,ok\n0008," + b"x" * (64 * 1024 + 1) + b"\n"
    elif extension == "xlsx":
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", b"x" * (32 * 1024 * 1024 + 1))
        content = output.getvalue()
    else:
        content = parquet_bytes({"Value": ["x" * (9 * 1024 * 1024)]})
    source = tmp_path / f"oversize.{extension}"
    source.write_bytes(content)
    before = sidecar_http_client.get("/health").json()
    events = read_events(sidecar_http_client.post("/ingest-path-stream", json={"file_path": str(source)}))
    assert events[-1]["type"] == "error"
    assert not any(event["type"] == "done" or event.get("step_code") == "embedding" for event in events)
    after = sidecar_http_client.get("/health").json()
    assert (after["documents_count"], after["chunks_count"]) == (before["documents_count"], before["chunks_count"])
    assert source.read_bytes() == content


def workbook_bytes(rows, extra_sheets=()):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    for name, sheet_rows in extra_sheets:
        sheet = book.create_sheet(name)
        for row in sheet_rows:
            sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


@pytest.fixture(params=["xlsx", "xls"])
def excel_source(request, tmp_path):
    path = tmp_path / f"cells.{request.param}"
    if request.param == "xlsx":
        content = workbook_bytes([
            ["Code", "Amount", "Note"],
            ["00123", 42.5, "NA"],
            ["AB123", 0, "ultimo valore"],
        ], [("Second", [["Final cell 9988"]])])
        expected = ["00123", "42.5", "NA", "AB123", "ultimo valore", "Final cell 9988"]
    else:
        content = (FIXTURES / "ragged.xls").read_bytes()
        expected = ["| a", "| d", "| f", "| g", "| k", "| j", "| l"]
    path.write_bytes(content)
    return path, expected


@pytest.mark.parametrize("from_disk", [False, True])
def test_excel_retains_cells(excel_source, from_disk):
    path, expected = excel_source
    original = path.read_bytes()
    markdown, pages = extract_document_markdown(
        path.name, original if not from_disk else b"", str(path) if from_disk else None,
    )
    assert pages == 1
    for value in expected:
        assert value in markdown
    assert path.read_bytes() == original


@pytest.mark.parametrize("rows", [[], [["   ", "\t"]]])
def test_excel_rejects_empty_cells(rows):
    with pytest.raises(ValueError, match="no readable cells"):
        extract_document_markdown("empty.xlsx", workbook_bytes(rows))


def test_excel_preserves_a_single_source_cell():
    markdown, _ = extract_document_markdown("single.xlsx", workbook_bytes([["AB123"]]))
    assert "AB123" in markdown


def test_excel_rejects_an_empty_selected_row_range():
    with pytest.raises(ValueError, match="no readable cells"):
        extract_document_markdown(
            "limited.xlsx", workbook_bytes([["   "], ["AB123"]]), max_excel_rows=1,
        )


def test_excel_reports_sheet_and_row_limits():
    content = workbook_bytes([["first"], ["second"], ["third"]], [("Other", [["hidden"]])])
    markdown, _ = extract_document_markdown(
        "limited.xlsx", content, max_excel_rows=2, max_sheets=1,
    )
    assert "first" in markdown and "second" in markdown
    assert "third" not in markdown and "hidden" not in markdown
    assert "su 3 totali" in markdown and "fogli su 2 totali" in markdown


def test_xlsx_counts_omitted_rows_beyond_first_batch():
    content = workbook_bytes([[f"row-{index}"] for index in range(1100)])
    markdown, _ = extract_document_markdown("many.xlsx", content, max_excel_rows=2)
    assert "su 1100 totali" in markdown
    assert "row-1099" not in markdown


def test_xlsx_reports_selected_sheets_without_readable_cells():
    content = workbook_bytes([], [("Second", [["visible"]])])
    markdown, _ = extract_document_markdown("sheets.xlsx", content)
    assert "visible" in markdown
    assert "Fogli selezionati senza celle leggibili nel limite di righe: 1" in markdown


@pytest.mark.parametrize("from_disk", [False, True])
def test_xls_oversize_is_refused_before_parsing(tmp_path, from_disk):
    content = b"x" * (8 * 1024 * 1024 + 1)
    path = tmp_path / "oversize.xls"
    if from_disk:
        path.write_bytes(content)
    with pytest.raises(ValueError, match="8 MiB parser limit"):
        extract_document_markdown(path.name, b"" if from_disk else content, str(path) if from_disk else None)


def test_xlsx_oversize_shared_strings_are_refused_before_loading():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", b"x" * (16 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="16 MiB parser limit"):
        extract_document_markdown("oversize.xlsx", output.getvalue())


def test_xlsx_refuses_compressed_worksheet_before_loading_and_large_inline_cell():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"x" * (32 * 1024 * 1024 + 1))
    assert len(output.getvalue()) < 100_000
    with pytest.raises(ValueError, match="XML payload exceeds"):
        extract_document_markdown("compressed.xlsx", output.getvalue(), allow_full_excel_format=True)

    original = workbook_bytes([["seed"]])
    inline = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(inline, "w") as target:
        for entry in source.infolist():
            payload = source.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                assert b"<t>seed</t>" in payload
                payload = payload.replace(b"<t>seed</t>", b"<t>" + b"x" * (64 * 1024 + 1) + b"</t>")
            target.writestr(entry, payload)
    content = inline.getvalue()
    with pytest.raises(ValueError, match="cell exceeds"):
        extract_document_markdown("inline.xlsx", content)


def test_full_format_xlsx_keeps_last_coordinate_in_bounded_rag_extract():
    book = Workbook()
    book.active.cell(row=1_048_576, column=16_384, value="last-cell")
    output = io.BytesIO()
    book.save(output)
    markdown, _ = extract_document_markdown("limits.xlsx", output.getvalue(), allow_full_excel_format=True)
    assert "1048576" in markdown
    assert "Col 16384" in markdown
    assert "last-cell" in markdown
    assert "Estratto RAG" in markdown


def test_full_format_xlsx_declares_row_column_and_cell_selection(monkeypatch):
    from sidecar.domain import ingestion

    book = Workbook()
    sheet = book.active
    for column in range(1, 130):
        sheet.cell(row=1, column=column, value="x" * 300 if column == 1 else f"value-{column}")
    for row in range(2, 152):
        sheet.cell(row=row, column=1, value=f"row-{row}")
    output = io.BytesIO()
    book.save(output)
    monkeypatch.setattr(ingestion, "XLSX_MAX_SHARED_STRINGS_BYTES", -1)
    with pytest.raises(ValueError, match="parser limit"):
        extract_document_markdown("selection.xlsx", output.getvalue())
    markdown, _ = extract_document_markdown("selection.xlsx", output.getvalue(), allow_full_excel_format=True)
    assert "prime 150 righe non vuote su 151" in markdown
    assert "prime 128 colonne non vuote su 129" in markdown
    assert "Colonne omesse dall'estratto: 1" in markdown
    assert "Celle abbreviate a 256 caratteri: 1" in markdown
    assert "row-151" not in markdown


def test_full_format_xls_bypasses_default_byte_gate_but_keeps_coordinates(monkeypatch):
    from sidecar.domain import ingestion

    content = (FIXTURES / "ragged.xls").read_bytes()
    monkeypatch.setattr(ingestion, "XLS_MAX_SOURCE_BYTES", len(content) - 1)
    with pytest.raises(ValueError, match="parser limit"):
        extract_document_markdown("ragged.xls", content)
    markdown, _ = extract_document_markdown("ragged.xls", content, allow_full_excel_format=True)
    assert "Estratto RAG" in markdown
    assert "Col 1" in markdown


@pytest.mark.parametrize("extension,engine", [("xlsx", "openpyxl"), ("xls", "xlrd")])
def test_excel_missing_engine_is_an_error(monkeypatch, extension, engine):
    content = workbook_bytes([["AB123"]]) if extension == "xlsx" else (FIXTURES / "ragged.xls").read_bytes()
    monkeypatch.setitem(sys.modules, engine, None)
    with pytest.raises(ImportError, match=engine):
        extract_document_markdown(f"missing.{extension}", content)


def test_excel_stream_retains_cells(sidecar_http_client, excel_source):
    path, expected = excel_source
    original = path.read_bytes()
    payload = done_payload(sidecar_http_client.post("/ingest-path-stream", json={"file_path": str(path)}))
    assert payload["num_chunks"] > 0
    for value in expected:
        assert value in payload["extracted_markdown"]
    assert path.read_bytes() == original
    assert sidecar_http_client.delete(f"/documents/{payload['id']}").status_code == 200
    assert path.read_bytes() == original


@pytest.mark.parametrize("name,content", [
    ("corrupt.xlsx", b"not an Excel workbook"),
    ("corrupt.xls", b"not an Excel workbook"),
    ("empty.xlsx", workbook_bytes([])),
    ("empty.xls", b""),
])
def test_excel_stream_failure_cannot_index(sidecar_http_client, tmp_path, name, content):
    path = tmp_path / name
    path.write_bytes(content)
    before = sidecar_http_client.get("/health").json()
    events = read_events(sidecar_http_client.post("/ingest-path-stream", json={"file_path": str(path)}))
    assert events[-1]["type"] == "error"
    assert events[-1]["error"]
    assert not any(event["type"] == "done" for event in events)
    assert not any(event.get("step_code") == "embedding" for event in events)
    after = sidecar_http_client.get("/health").json()
    assert after["documents_count"] == before["documents_count"]
    assert after["chunks_count"] == before["chunks_count"]
    assert path.read_bytes() == content


def parquet_bytes(data):
    output = io.BytesIO()
    pd.DataFrame(data).to_parquet(output, engine="pyarrow", index=False)
    return output.getvalue()


@pytest.mark.parametrize("from_disk", [False, True])
def test_parquet_retains_values(tmp_path, from_disk):
    content = parquet_bytes({"code": ["00123", "AB123"], "amount": [42.5, 0], "note": ["NA", "ultimo valore"]})
    path = tmp_path / "cells.parquet"
    path.write_bytes(content)
    markdown, pages = extract_document_markdown(
        path.name, content if not from_disk else b"", str(path) if from_disk else None,
    )
    assert pages == 1
    for value in ("code", "amount", "note", "00123", "AB123", "42.5", "NA", "ultimo valore"):
        assert value in markdown
    assert path.read_bytes() == content


def test_parquet_reports_row_limit():
    content = parquet_bytes({"code": ["first", "second", "third"]})
    markdown, _ = extract_document_markdown("limited.parquet", content, max_tabular_rows=2)
    assert "first" in markdown and "second" in markdown
    assert "third" not in markdown and "su 3 totali" in markdown


def test_parquet_counts_omitted_rows_across_batches():
    content = parquet_bytes({"code": [f"row-{index}" for index in range(1100)]})
    markdown, _ = extract_document_markdown("many.parquet", content, max_tabular_rows=2)
    assert "su 1100 totali" in markdown
    assert "row-1099" not in markdown


def test_parquet_refuses_wide_schema_and_oversize_decompressed_batch():
    wide = parquet_bytes({f"C{index}": ["x"] for index in range(257)})
    with pytest.raises(ValueError, match="column parser limit"):
        extract_document_markdown("wide.parquet", wide)
    long_header = parquet_bytes({"H" * (64 * 1024 + 1): ["x"]})
    with pytest.raises(ValueError, match="cell exceeds"):
        extract_document_markdown("long-header.parquet", long_header)
    compressed = parquet_bytes({"Value": ["x" * (9 * 1024 * 1024)]})
    assert len(compressed) < 1024 * 1024
    with pytest.raises(ValueError, match="batch exceeds"):
        extract_document_markdown("compressed.parquet", compressed)


def test_parquet_rejects_an_empty_selected_row_range():
    with pytest.raises(ValueError, match="no readable values"):
        extract_document_markdown(
            "limited.parquet", parquet_bytes({"code": [None, "AB123"]}), max_tabular_rows=1,
        )


@pytest.mark.parametrize("data", [{"code": []}, {"code": [None]}, {"code": ["   "]}])
def test_parquet_rejects_empty_values(data):
    with pytest.raises(ValueError, match="no readable values"):
        extract_document_markdown("empty.parquet", parquet_bytes(data))


def test_parquet_missing_engine_is_an_error(monkeypatch):
    content = parquet_bytes({"code": ["AB123"]})
    monkeypatch.setitem(sys.modules, "pyarrow", None)
    with pytest.raises(ImportError, match="pyarrow"):
        extract_document_markdown("missing.parquet", content)


def test_parquet_stream_retains_values(sidecar_http_client, tmp_path):
    content = parquet_bytes({"code": ["00123", "AB123"], "amount": [42.5, 0], "note": ["NA", "ultimo valore"]})
    path = tmp_path / "cells.parquet"
    path.write_bytes(content)
    payload = done_payload(sidecar_http_client.post("/ingest-path-stream", json={"file_path": str(path)}))
    assert payload["num_chunks"] > 0
    for value in ("00123", "AB123", "42.5", "NA", "ultimo valore"):
        assert value in payload["extracted_markdown"]
    assert path.read_bytes() == content
    assert sidecar_http_client.delete(f"/documents/{payload['id']}").status_code == 200
    assert path.read_bytes() == content


@pytest.mark.parametrize("content", [b"not Parquet", b"", parquet_bytes({"code": []})])
def test_parquet_stream_failure_cannot_index(sidecar_http_client, tmp_path, content):
    path = tmp_path / "invalid.parquet"
    path.write_bytes(content)
    before = sidecar_http_client.get("/health").json()
    events = read_events(sidecar_http_client.post("/ingest-path-stream", json={"file_path": str(path)}))
    assert events[-1]["type"] == "error" and events[-1]["error"]
    assert not any(event["type"] == "done" for event in events)
    assert not any(event.get("step_code") == "embedding" for event in events)
    after = sidecar_http_client.get("/health").json()
    assert after["documents_count"] == before["documents_count"]
    assert after["chunks_count"] == before["chunks_count"]
    assert path.read_bytes() == content
