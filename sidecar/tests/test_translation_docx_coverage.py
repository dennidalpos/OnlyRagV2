"""Native package readback independent of the production DOCX collector."""
from collections import Counter
import json
import io
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import docx
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.shared import Inches
from PIL import Image
import pytest

from sidecar.domain import translator


def package_text(path):
    with ZipFile(path) as package:
        return {
            name: [element.text for element in ET.fromstring(package.read(name)).iter(qn("w:t")) if element.text]
            for name in package.namelist()
            if name == "word/document.xml" or name.startswith(("word/header", "word/footer")) and name.endswith(".xml")
        }


def record(source):
    return dict(id="coverage", file_path=str(source), filename=source.name, file_type="docx")


def test_native_docx_translates_all_story_parts_nested_merged_and_hyperlink_text_once(tmp_path, monkeypatch):
    source = tmp_path / "stories.docx"
    document = docx.Document()
    document.add_paragraph().add_run("SOURCE body* user@example.test").bold = True
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "SOURCE merged"
    nested = merged.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "SOURCE nested"
    image = io.BytesIO()
    Image.new("RGB", (5, 5), "red").save(image, format="PNG")
    document.add_paragraph().add_run().add_picture(io.BytesIO(image.getvalue()))
    document.add_paragraph("SOURCE after table")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), "bookmark")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "SOURCE hyperlink"
    run.append(text)
    hyperlink.append(run)
    document.add_paragraph()._p.append(hyperlink)
    section = document.sections[0]
    section.different_first_page_header_footer = True
    document.settings.odd_and_even_pages_header_footer = True
    for name in ("header", "footer", "first_page_header", "first_page_footer", "even_page_header", "even_page_footer"):
        story = getattr(section, name)
        story.paragraphs[0].add_run(f"SOURCE {name}*").italic = True
        story.add_table(rows=1, cols=1, width=Inches(3)).cell(0, 0).text = f"SOURCE {name} table"
    second = document.add_section(WD_SECTION_START.NEW_PAGE)
    assert second.header.is_linked_to_previous
    second.footer.is_linked_to_previous = False
    second.footer.paragraphs[0].text = "SOURCE second footer"
    document.save(source)
    original = source.read_bytes()
    before = package_text(source)
    expected = [text for texts in before.values() for text in texts]
    loaded = docx.Document(source)
    collected = translator._collect_docx_runs(loaded)
    assert Counter(run.text for run in collected) == Counter(expected)
    assert [run.text for run in collected[:4]] == ["SOURCE body* user@example.test", "SOURCE merged", "SOURCE nested", "SOURCE after table"]
    calls = []

    def fake_transport(text, *args, **kwargs):
        calls.append(text)
        return text.replace("SOURCE", "TARGET")

    monkeypatch.setattr(translator, "_call_ollama_translate", fake_transport)
    result = [json.loads(line) for line in translator.translate_document_stream(record(source), "English", "Italian", "transport-fixture")]
    assert result[-1]["type"] == "done", result
    output = result[-1]["data"]["filePath"]
    after = package_text(output)
    assert after == {part: [text.replace("SOURCE", "TARGET") for text in texts] for part, texts in before.items()}
    assert all(text.count("SOURCE") == 0 for texts in after.values() for text in texts)
    assert sum(text.count("*") for texts in after.values() for text in texts) == 7
    assert docx.Document(output).paragraphs[0].runs[0].bold is True
    assert docx.Document(output).sections[0].header.paragraphs[0].runs[0].italic is True
    with ZipFile(source) as old, ZipFile(output) as new:
        for name in old.namelist():
            if name not in before:
                assert old.read(name) == new.read(name), name
    assert calls
    assert source.read_bytes() == original


@pytest.mark.parametrize("unsupported", ["txbxContent", "altChunk", "fldSimple", "fldChar", "instrText", "footnoteReference", "endnoteReference", "sdt", "sym", "ins", "del"])
def test_unsupported_docx_content_is_refused_before_transport(tmp_path, monkeypatch, unsupported):
    source = tmp_path / "unsupported.docx"
    document = docx.Document()
    paragraph = document.add_paragraph("SOURCE body")
    paragraph._p.append(OxmlElement(f"w:{unsupported}"))
    document.save(source)
    original = source.read_bytes()
    def forbidden(*args, **kwargs):
        pytest.fail("Unsupported visible content must be refused before model transport")
    monkeypatch.setattr(translator, "_call_ollama_translate", forbidden)
    result = [json.loads(line) for line in translator.translate_document_stream(record(source), "English", "Italian", "transport-fixture")]
    assert result[-1]["type"] == "error"
    assert "Unsupported DOCX" in result[-1]["error"]
    assert not any(event["type"] == "done" for event in result)
    assert list(tmp_path.iterdir()) == [source]
    assert source.read_bytes() == original


def test_collection_does_not_materialize_unset_header_footer_parts():
    document = docx.Document()
    document.add_paragraph("SOURCE body")
    before = set(document.part.rels)
    assert [run.text for run in translator._collect_docx_runs(document)] == ["SOURCE body"]
    assert set(document.part.rels) == before


def test_mixed_text_and_graphic_run_is_refused_instead_of_erasing_graphic():
    document = docx.Document()
    run = document.add_paragraph().add_run("SOURCE text")
    run._r.append(OxmlElement("w:drawing"))
    with pytest.raises(ValueError, match="mixed text/graphic"):
        translator._collect_docx_runs(document)


def test_text_in_an_additional_word_part_is_refused(tmp_path):
    source = tmp_path / "footnotes.docx"
    document = docx.Document()
    document.add_paragraph("SOURCE body")
    part = XmlPart(PackURI("/word/footnotes.xml"), "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                   parse_xml('<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="1"><w:p><w:r><w:t>SOURCE note</w:t></w:r></w:p></w:footnote></w:footnotes>'), document.part.package)
    document.part.relate_to(part, RT.FOOTNOTES)
    document.save(source)
    with pytest.raises(ValueError, match="Unsupported DOCX text part: /word/footnotes.xml"):
        translator._collect_docx_runs(docx.Document(source))


@pytest.mark.parametrize("tag", ["a:t", "c:chart", "dgm:relIds", "m:oMath"])
def test_other_visible_text_containers_are_refused(tag):
    document = docx.Document()
    document.add_paragraph("SOURCE body")._p.append(OxmlElement(tag))
    with pytest.raises(ValueError, match="Unsupported DOCX content"):
        translator._collect_docx_runs(document)
