"""Native PDF/DOCX publication, with deterministic translation transport only."""
import json
from pathlib import Path
from zipfile import ZipFile

import docx
import pymupdf
import pytest

from sidecar.domain import translator
from sidecar.infrastructure.translation_output import TranslationOutput
from sidecar.services.task_cancellation import cancel_task
from _stream import retained_translation_record


@pytest.fixture(params=["docx", "pdf"])
def translation(request, tmp_path, monkeypatch):
    kind = request.param
    source = tmp_path / f"source.{kind}"
    if kind == "docx":
        document = docx.Document()
        document.add_paragraph("Source text")
        document.save(source)
    else:
        with pymupdf.open() as document:
            document.new_page().insert_text((72, 72), "Source text")
            document.save(source)

    def translate_runs(runs, *args):
        for run in runs:
            run.text = "Tradotto"

    def translate_blocks(blocks, *args):
        for block in blocks:
            block["text"] = "Tradotto"

    monkeypatch.setattr(translator, "_translate_batch", translate_runs)
    monkeypatch.setattr(translator, "_translate_pdf_blocks", translate_blocks)
    record = retained_translation_record(source)
    original = source.read_bytes()

    def stream(task_id=None):
        return translator.translate_document_stream(record, "English", "Italian", "transport-fixture", task_id=task_id)

    yield kind, source, tmp_path / f"source_italian.{kind}", stream
    assert source.read_bytes() == original


def events(stream):
    return [json.loads(line) for line in stream]


def assert_complete(path, kind):
    if kind == "docx":
        assert docx.Document(path).paragraphs[0].text == "Tradotto"
        with ZipFile(path) as package:
            assert package.testzip() is None
    else:
        with pymupdf.open(path) as document:
            assert "Tradotto" in document[0].get_text()


def test_publication_is_complete_and_preserves_prior_collision(translation):
    kind, source, destination, stream = translation
    destination.write_bytes(b"unrelated existing bytes")
    result = events(stream())
    assert result[-1]["type"] == "done"
    output = Path(result[-1]["data"]["filePath"])
    assert output != destination
    assert_complete(output, kind)
    assert result[-1]["data"]["file_size"] == output.stat().st_size
    assert destination.read_bytes() == b"unrelated existing bytes"
    assert not list(source.parent.glob(".onlyrag-translation-*"))


def test_late_collision_is_exclusive_and_a_later_attempt_recovers(translation):
    kind, source, destination, stream = translation
    pending = stream()
    assert json.loads(next(pending))["type"] == "start"
    destination.write_bytes(b"created during translation")
    assert events(pending)[-1]["type"] == "error"
    assert destination.read_bytes() == b"created during translation"
    assert not list(source.parent.glob(".onlyrag-translation-*"))
    recovered = events(stream())[-1]
    assert recovered["type"] == "done"
    assert_complete(recovered["data"]["filePath"], kind)


def test_save_failure_after_actual_writes_never_exposes_partial_final(translation, monkeypatch):
    kind, source, destination, stream = translation
    if kind == "docx":
        original = ZipFile.writestr
        written = []

        def failing_write(package, *args, **kwargs):
            result = original(package, *args, **kwargs)
            written.append(1)
            if len(written) == 2:
                raise OSError("injected ZIP write failure")
            return result

        monkeypatch.setattr(ZipFile, "writestr", failing_write)
    else:
        original = pymupdf.Document.save

        def failing_save(document, path, *args, **kwargs):
            original(document, path, *args, **kwargs)
            assert Path(path).stat().st_size > 0
            raise OSError("injected PDF save failure after writes")

        monkeypatch.setattr(pymupdf.Document, "save", failing_save)
    result = events(stream())
    assert result[-1]["type"] == "error"
    assert not destination.exists()
    assert list(source.parent.iterdir()) == [source]
    monkeypatch.undo()


def test_cancellation_after_serialization_discards_staging(translation, monkeypatch):
    _, source, destination, stream = translation
    original = TranslationOutput.validate

    def cancel_after_validation(output, pages):
        size = original(output, pages)
        cancel_task("cancel-publication")
        return size

    monkeypatch.setattr(TranslationOutput, "validate", cancel_after_validation)
    result = events(stream("cancel-publication"))
    assert result[-1] == dict(type="cancelled", task_id="cancel-publication")
    assert not destination.exists()
    assert list(source.parent.iterdir()) == [source]


@pytest.mark.parametrize("after_publish", [False, True])
def test_cleanup_refusal_reports_recoverable_paths(translation, monkeypatch, after_publish):
    kind, source, destination, stream = translation
    original = Path.unlink

    def refuse_staging(path, *args, **kwargs):
        if path.parent.name.startswith(".onlyrag-translation-"):
            raise PermissionError("injected cleanup refusal")
        return original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "unlink", refuse_staging)
    if not after_publish:
        def fail_publication(output):
            raise PermissionError("injected publication refusal")
        monkeypatch.setattr(TranslationOutput, "publish", fail_publication)
    result = events(stream())
    assert result[-1]["type"] == "error"
    assert "cleanup failed" in result[-1]["error"]
    stage = next(source.parent.glob(".onlyrag-translation-*"))
    assert str(stage) in result[-1]["error"]
    assert destination.exists() == after_publish
    if after_publish:
        assert str(destination) in result[-1]["error"]
        assert_complete(destination, kind)
    assert_complete(stage / f"output.{kind}", kind)
    monkeypatch.undo()
    (stage / f"output.{kind}").unlink()
    stage.rmdir()


def test_generator_close_reclaims_staging(translation):
    _, source, destination, stream = translation
    pending = stream()
    assert json.loads(next(pending))["type"] == "start"
    pending.close()
    assert not destination.exists()
    assert list(source.parent.iterdir()) == [source]


def test_generator_close_cleanup_refusal_propagates_and_unregisters_task(translation, monkeypatch):
    _, source, destination, stream = translation
    from sidecar.services import task_cancellation
    pending = stream("closed-cleanup-refusal")
    next(pending)
    original = Path.rmdir
    def refuse_directory(path, *args, **kwargs):
        if path.name.startswith(".onlyrag-translation-"):
            raise PermissionError("injected close cleanup refusal")
        return original(path, *args, **kwargs)
    monkeypatch.setattr(Path, "rmdir", refuse_directory)
    with pytest.raises(OSError, match="staging cleanup failed"):
        pending.close()
    assert "closed-cleanup-refusal" not in task_cancellation._tasks
    assert not destination.exists()
    stage = next(source.parent.glob(".onlyrag-translation-*"))
    monkeypatch.undo()
    stage.rmdir()


def test_overlapping_translations_cannot_replace_the_winner(translation):
    kind, source, destination, stream = translation
    first, second = stream(), stream()
    next(first)
    next(second)
    assert events(first)[-1]["type"] == "done"
    completed = destination.read_bytes()
    assert events(second)[-1]["type"] == "error"
    assert destination.read_bytes() == completed
    assert_complete(destination, kind)
    assert not list(source.parent.glob(".onlyrag-translation-*"))


def test_invalid_serialized_output_is_not_published(translation, monkeypatch):
    kind, source, destination, stream = translation
    def corrupt_save(document, output, *args, **kwargs):
        if hasattr(output, "write"):
            output.write(b"not a valid document")
        else:
            Path(output).write_bytes(b"not a valid document")
    monkeypatch.setattr(docx.document.Document if kind == "docx" else pymupdf.Document, "save", corrupt_save)
    assert events(stream())[-1]["type"] == "error"
    assert not destination.exists()
    assert list(source.parent.iterdir()) == [source]


def test_unsupported_publication_filesystem_fails_without_copy_fallback(translation, monkeypatch):
    _, source, destination, stream = translation
    def unsupported_link(*args, **kwargs):
        raise OSError("hard links unavailable on this filesystem")
    monkeypatch.setattr(translator.os, "link", unsupported_link)
    assert "hard links unavailable" in events(stream())[-1]["error"]
    assert not destination.exists()
    assert list(source.parent.iterdir()) == [source]


@pytest.mark.parametrize("filename,language", [("../escape.docx", "Italian"), ("source.docx", "../Italian"), ("C:\\escape.docx", "Italian")])
def test_invalid_output_components_are_refused(tmp_path, filename, language):
    with pytest.raises(ValueError, match="Invalid translation"):
        translator._resolve_output_filepath(str(tmp_path / "source.docx"), filename, language)
